"""Small adapters around the actual libraries, with honest source provenance."""
from __future__ import annotations

import io
import re
import tempfile
import time
from dataclasses import dataclass, field, asdict
from functools import lru_cache
from pathlib import Path
from typing import Any


TOOLS = {
    "Baseline": ["Paragraphs", "Pages", "Fixed characters", "Markdown headings"],
    "LangChain": ["Recursive characters", "Recursive tokens", "Fixed tokens", "Markdown headings"],
    "LlamaIndex": ["Sentence", "Tokens", "Markdown headings", "Hierarchical"],
    "Chonkie": ["Tokens", "Sentence", "Recursive", "Semantic (local model)"],
    "Haystack": ["Words", "Paragraphs", "Pages", "Sentence", "Recursive"],
    "Docling": ["Hierarchical", "Hybrid"],
    "Unstructured": ["Basic elements", "By title", "By title + page boundaries"],
    "semchunk": ["Recursive tokens"],
    "semantic-text-splitter": ["Text structure", "Markdown structure"],
}


def strategy_info(tool: str, strategy: str) -> tuple[str, str, bool]:
    if tool == "Baseline":
        return {
            "Paragraphs": ("paragraphs", "Groups whole detected paragraphs; does not impose a token limit.", True),
            "Pages": ("pages", "Uses PDF pages or explicit page-break groups. Never invents rendered DOCX pages.", True),
            "Fixed characters": ("characters", "Cuts at a character count, even in the middle of a sentence.", True),
            "Markdown headings": ("none", "One chunk per Markdown heading section; large sections stay large.", False),
        }[strategy]
    if tool == "LangChain":
        return {
            "Recursive characters": ("characters", "Tries paragraph, newline, word, then character boundaries.", True),
            "Recursive tokens": ("tokens", "Recursive splitting, measured with cl100k_base tokens.", True),
            "Fixed tokens": ("tokens", "Fixed tiktoken windows; may split a Unicode character across chunks.", True),
            "Markdown headings": ("none", "Native MarkdownHeaderTextSplitter; headings are retained.", False),
        }[strategy]
    if tool == "LlamaIndex":
        return {
            "Sentence": ("tokens", "Native SentenceSplitter tries to preserve complete sentences.", True),
            "Tokens": ("tokens", "Native TokenTextSplitter uses token budgets and text separators.", True),
            "Markdown headings": ("none", "Native MarkdownNodeParser splits by heading sections.", False),
            "Hierarchical": ("tokens", "Creates parent and child chunks at 4x, 2x and 1x the chosen size.", True),
        }[strategy]
    if tool == "Chonkie":
        return ("tokens", {
            "Tokens": "Fixed token chunks using cl100k_base.",
            "Sentence": "Groups sentences into token-budget chunks.",
            "Recursive": "Uses native recursive boundary rules. Overlap is not supported here.",
            "Semantic (local model)": "Finds topic changes with a local MiniLM embedding model. First use downloads model weights.",
        }[strategy], strategy in {"Tokens", "Sentence"})
    if tool == "Haystack":
        units = {"Words": "words", "Paragraphs": "paragraphs", "Pages": "pages", "Sentence": "sentences", "Recursive": "tokens"}
        return units[strategy], "Native Haystack splitter; size and overlap count the selected units.", True
    if tool == "Docling":
        return ("tokens" if strategy == "Hybrid" else "none", "Native DoclingDocument chunking. Hybrid splits oversized elements and merges compatible small elements; hierarchical follows document elements.", False)
    if tool == "Unstructured":
        return "characters", "Native partitioning + chunking. Tables stay separate. Overlap applies to oversized elements; title strategies preserve sections.", True
    return ("tokens", "Native boundary-aware splitter with a cl100k_base token budget; no embedding model is used.", True)


@dataclass
class Block:
    text: str
    page: int | None = None
    kind: str = "paragraph"
    start: int = 0
    end: int = 0


@dataclass
class Parsed:
    name: str
    parser: str
    text: str
    blocks: list[Block]
    page_basis: str
    notes: list[str] = field(default_factory=list)
    native: Any = None
    elapsed: float = 0


@dataclass
class Chunk:
    id: int
    text: str
    tokens: int
    characters: int
    pages: list[int]
    blocks: list[int]
    start: int | None
    end: int | None
    metadata: dict = field(default_factory=dict)


@lru_cache(maxsize=1)
def encoding():
    import tiktoken
    return tiktoken.get_encoding("cl100k_base")


def token_ids(text: str):
    return encoding().encode(text, disallowed_special=())


def assemble(name, parser, blocks, basis, notes=None, native=None):
    parts, offset, prev_page = [], 0, None
    for block in blocks:
        # Form feeds provide an explicit boundary to Haystack's page splitter.
        separator = "\n\n\f\n\n" if parts and block.page != prev_page and block.page is not None else "\n\n" if parts else ""
        parts.append(separator)
        offset += len(separator)
        block.start = offset
        parts.append(block.text)
        offset += len(block.text)
        block.end = offset
        prev_page = block.page
    return Parsed(name, parser, "".join(parts), blocks, basis, notes or [], native)


def markdown_blocks(text):
    blocks, page = [], 1
    has_pages = bool(re.search(r"\f|<!--\s*pagebreak\s*-->", text, re.I))
    for segment in re.split(r"\f|<!--\s*pagebreak\s*-->", text, flags=re.I):
        # Keep fenced code intact when blank lines occur inside it.
        pieces, current, fence = [], [], None
        for line in segment.splitlines():
            match = re.match(r"^\s*(`{3,}|~{3,})", line)
            if match:
                marker = match.group(1)[0]
                fence = None if fence == marker else marker if fence is None else fence
            if not line.strip() and fence is None:
                if current:
                    pieces.append("\n".join(current))
                    current = []
            else:
                current.append(line)
        if current:
            pieces.append("\n".join(current))
        for piece in pieces:
            kind = "heading" if re.match(r"^#{1,6}\s", piece) else "code" if re.match(r"^\s*(```|~~~)", piece) else "table" if re.search(r"\|.*\|", piece) else "paragraph"
            blocks.append(Block(piece, page if has_pages else None, kind))
        page += 1
    return blocks, "Explicit Markdown page-break groups" if has_pages else "Markdown has no physical pages"


def quick_parse(name: str, data: bytes) -> Parsed:
    suffix = Path(name).suffix.lower()
    if suffix in {".md", ".markdown"}:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("Save the Markdown file as UTF-8 and try again.") from exc
        blocks, basis = markdown_blocks(text.replace("\r\n", "\n"))
        return assemble(name, "Quick text", blocks, basis)
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ValueError("This PDF needs a password. Upload an unlocked copy.")
        blocks, empty = [], []
        for page_no, page in enumerate(reader.pages, 1):
            text = (page.extract_text(extraction_mode="layout") or "").strip()
            if not text:
                empty.append(page_no)
                blocks.append(Block("", page_no, "empty page"))
            for part in re.split(r"\n\s*\n", text):
                if part.strip():
                    blocks.append(Block(part.strip(), page_no))
        notes = ["PDF paragraph boundaries are inferred from extracted whitespace; layout or OCR parsers can produce different boundaries."]
        if empty:
            notes.append(f"No extractable text on PDF pages {empty}. Select Docling with OCR for scanned pages.")
        return assemble(name, "Quick text", blocks, "Physical PDF pages", notes)
    if suffix == ".docx":
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        doc, blocks, page, breaks = Document(io.BytesIO(data)), [], 1, 0
        for element in doc.element.body.iterchildren():
            if element.tag == qn("w:p"):
                paragraph = Paragraph(element, doc)
                if paragraph.paragraph_format.page_break_before and blocks:
                    page += 1
                    breaks += 1
                kind = "heading" if paragraph.style and paragraph.style.name.startswith("Heading") else "paragraph"
                buffer = ""
                for node in element.iter():
                    if node.tag == qn("w:t"):
                        buffer += node.text or ""
                    elif node.tag == qn("w:tab"):
                        buffer += "\t"
                    elif node.tag == qn("w:br"):
                        if node.get(qn("w:type")) == "page":
                            if buffer.strip():
                                blocks.append(Block(buffer, page, kind))
                            buffer = ""
                            page += 1
                            breaks += 1
                        else:
                            buffer += "\n"
                if buffer.strip():
                    if kind == "heading":
                        level_match = re.search(r"(\d+)", paragraph.style.name)
                        level = min(int(level_match.group(1)), 6) if level_match else 1
                        buffer = "#" * level + " " + buffer
                    blocks.append(Block(buffer, page, kind))
            elif element.tag == qn("w:tbl"):
                table = Table(element, doc)
                rows = ["| " + " | ".join(cell.text.replace("\n", " ").replace("|", "\\|") for cell in row.cells) + " |" for row in table.rows]
                if rows:
                    rows.insert(1, "| " + " | ".join("---" for _ in table.rows[0].cells) + " |")
                    blocks.append(Block("\n".join(rows), page, "table"))
        if not breaks:
            for block in blocks:
                block.page = None
        return assemble(name, "Quick text", blocks, "DOCX explicit page-break groups (not rendered pages)" if breaks else "DOCX has no explicit page breaks; rendered pages are unknown", ["DOCX pagination depends on Word's layout. This demo uses explicit page breaks only; it ignores cached rendered breaks."])
    raise ValueError("Choose a PDF, DOCX or Markdown file.")


def parse_document(name: str, data: bytes, parser="Quick text", ocr=False) -> Parsed:
    started = time.perf_counter()
    if parser == "Quick text":
        parsed = quick_parse(name, data)
    else:
        suffix = Path(name).suffix.lower()
        with tempfile.TemporaryDirectory(prefix="chunk_lab_") as folder:
            path = Path(folder) / ("document" + suffix)
            path.write_bytes(data)
            if parser == "Docling":
                from docling.document_converter import DocumentConverter, PdfFormatOption
                from docling.datamodel.base_models import InputFormat
                from docling.datamodel.pipeline_options import PdfPipelineOptions
                options = PdfPipelineOptions(do_ocr=ocr, do_table_structure=True)
                converter = DocumentConverter(format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=options)})
                native = converter.convert(path).document
                blocks = []
                for item, _ in native.iterate_items():
                    text = getattr(item, "text", None)
                    if hasattr(item, "export_to_markdown"):
                        text = item.export_to_markdown(doc=native)
                    if text:
                        provenance = getattr(item, "prov", [])
                        page = provenance[0].page_no if provenance and suffix == ".pdf" else None
                        label = str(getattr(item, "label", "paragraph"))
                        if "section_header" in label:
                            text = "#" * min(getattr(item, "level", 1) or 1, 6) + " " + text
                        blocks.append(Block(text, page, label))
                parsed = assemble(name, parser, blocks, "Physical PDF pages from Docling provenance" if suffix == ".pdf" else "No physical page mapping for this format", native=native)
            elif parser == "Unstructured":
                if suffix == ".pdf":
                    from unstructured.partition.pdf import partition_pdf
                    native = partition_pdf(filename=str(path), strategy="fast", include_page_breaks=True)
                elif suffix == ".docx":
                    from unstructured.partition.docx import partition_docx
                    native = partition_docx(filename=str(path))
                else:
                    from unstructured.partition.md import partition_md
                    native = partition_md(filename=str(path))
                blocks = []
                for item in native:
                    if not str(item).strip():
                        continue
                    page = getattr(item.metadata, "page_number", None)
                    # DOCX metadata may use cached rendered breaks; don't label these physical pages.
                    kind = getattr(item, "category", "element")
                    text = str(item)
                    blocks.append(Block(text, page, kind))
                parsed = assemble(name, parser, blocks, "Physical PDF pages" if suffix == ".pdf" else "Parser-reported page markers (not verified rendered pages)" if suffix == ".docx" else "Markdown has no physical pages", ["Unstructured PDF uses its fast text parser here. Use Docling + OCR for image-only PDFs."], native)
            else:
                raise ValueError("Unknown parser")
    parsed.elapsed = round(time.perf_counter() - started, 3)
    if not parsed.text.strip():
        raise ValueError("No text was extracted. For a scanned PDF, select Docling and enable OCR.")
    return parsed


def normalized_with_map(text):
    result, positions = [], []
    for match in re.finditer(r"\s+|\S", text):
        result.append(" " if match.group().isspace() else match.group())
        positions.append(match.start())
    return "".join(result), positions


def attach_provenance(parsed: Parsed, raw: list[dict]) -> list[Chunk]:
    normalized, positions = normalized_with_map(parsed.text)
    chunks = []
    for item in raw:
        text = item["text"]
        if not text.strip():
            continue
        start, end = item.get("start"), item.get("end")
        if start is None or end is None or parsed.text[start:end] != text:
            target = re.sub(r"\s+", " ", text).strip()
            found = normalized.find(target) if target else -1
            # Repeated text cannot be located reliably without native source offsets.
            ambiguous = found >= 0 and normalized.find(target, found + 1) >= 0
            if found >= 0 and not ambiguous:
                start = positions[found]
                end = positions[found + len(target) - 1] + 1
            else:
                start = end = None
        block_ids = [i + 1 for i, block in enumerate(parsed.blocks) if start is not None and end is not None and block.end > start and block.start < end]
        pages = item.get("pages")
        if pages is None:
            pages = sorted({parsed.blocks[i - 1].page for i in block_ids if parsed.blocks[i - 1].page is not None})
        chunks.append(Chunk(len(chunks) + 1, text, len(token_ids(text)), len(text), pages, block_ids, start, end, item.get("metadata", {})))
    return chunks


def text_chunks(tool, strategy, text, size, overlap):
    if tool == "LangChain":
        from langchain_text_splitters import RecursiveCharacterTextSplitter, TokenTextSplitter, MarkdownHeaderTextSplitter
        if strategy == "Markdown headings":
            splitter = MarkdownHeaderTextSplitter(headers_to_split_on=[("#" * n, f"H{n}") for n in range(1, 7)], strip_headers=False)
            return [{"text": doc.page_content, "metadata": doc.metadata} for doc in splitter.split_text(text)]
        if strategy == "Fixed tokens":
            splitter = TokenTextSplitter(encoding_name="cl100k_base", chunk_size=size, chunk_overlap=overlap, disallowed_special=())
        elif strategy == "Recursive tokens":
            splitter = RecursiveCharacterTextSplitter.from_tiktoken_encoder(encoding_name="cl100k_base", chunk_size=size, chunk_overlap=overlap, disallowed_special=())
        else:
            splitter = RecursiveCharacterTextSplitter(chunk_size=size, chunk_overlap=overlap)
        return [{"text": value} for value in splitter.split_text(text)]
    if tool == "LlamaIndex":
        from llama_index.core import Document
        from llama_index.core.node_parser import SentenceSplitter, TokenTextSplitter, MarkdownNodeParser, HierarchicalNodeParser
        if strategy == "Markdown headings":
            splitter = MarkdownNodeParser()
        elif strategy == "Hierarchical":
            splitter = HierarchicalNodeParser.from_defaults(chunk_sizes=[size * 4, size * 2, size], chunk_overlap=overlap)
        else:
            splitter = (SentenceSplitter if strategy == "Sentence" else TokenTextSplitter)(chunk_size=size, chunk_overlap=overlap, tokenizer=token_ids)
        nodes = splitter.get_nodes_from_documents([Document(text=text)])
        raw = []
        for node in nodes:
            metadata = dict(node.metadata)
            metadata["node_id"] = node.node_id
            if node.parent_node:
                metadata["parent_id"] = node.parent_node.node_id
            if node.child_nodes:
                metadata["child_ids"] = [child.node_id for child in node.child_nodes]
            raw.append({"text": node.text, "metadata": metadata})
        return raw
    if tool == "Chonkie":
        from chonkie import TokenChunker, SentenceChunker, RecursiveChunker, SemanticChunker
        if strategy == "Semantic (local model)":
            splitter = SemanticChunker(embedding_model="sentence-transformers/all-MiniLM-L6-v2", chunk_size=size)
        elif strategy == "Recursive":
            splitter = RecursiveChunker(tokenizer=encoding(), chunk_size=size)
        else:
            splitter = (TokenChunker if strategy == "Tokens" else SentenceChunker)(tokenizer=encoding(), chunk_size=size, chunk_overlap=overlap)
        return [{"text": item.text, "start": item.start_index, "end": item.end_index, "metadata": {"native_token_count": item.token_count}} for item in splitter.chunk(text)]
    if tool == "Haystack":
        from haystack import Document
        from haystack.components.preprocessors import DocumentSplitter, RecursiveDocumentSplitter
        if strategy == "Recursive":
            splitter = RecursiveDocumentSplitter(split_length=size, split_overlap=overlap, split_unit="token")
        else:
            unit = {"Words": "word", "Paragraphs": "passage", "Pages": "page", "Sentence": "sentence"}[strategy]
            splitter = DocumentSplitter(split_by=unit, split_length=size, split_overlap=overlap, split_threshold=0)
        splitter.warm_up()
        documents = splitter.run(documents=[Document(content=text)])["documents"]
        return [{"text": doc.content, "metadata": doc.meta} for doc in documents]
    if tool == "semchunk":
        import semchunk
        splitter = semchunk.chunkerify(encoding(), size)
        values, offsets = splitter(text, overlap=overlap, offsets=True)
        return [{"text": value, "start": start, "end": end} for value, (start, end) in zip(values, offsets)]
    if tool == "semantic-text-splitter":
        from semantic_text_splitter import TextSplitter, MarkdownSplitter
        cls = MarkdownSplitter if strategy == "Markdown structure" else TextSplitter
        splitter = cls.from_tiktoken_model("gpt-4", size, overlap=overlap)
        return [{"text": value} for value in splitter.chunks(text)]
    raise ValueError(f"Unsupported tool: {tool}")


def run_chunking(parsed: Parsed, tool, strategy, size=256, overlap=0, keep_pages=False):
    started = time.perf_counter()
    if strategy not in TOOLS.get(tool, []):
        raise ValueError("Choose a supported tool and strategy.")
    units, _, supports_overlap = strategy_info(tool, strategy)
    if size < 1 or overlap < 0 or (supports_overlap and overlap >= size):
        raise ValueError("Size must be positive and overlap must be smaller than size.")
    if not supports_overlap:
        overlap = 0
    if (strategy == "Pages" or keep_pages) and not any(block.page is not None for block in parsed.blocks):
        raise ValueError("This document has no known page boundaries. Use a PDF or a DOCX with explicit page breaks.")
    notes, raw = [], []
    if tool == "Baseline":
        if strategy == "Fixed characters":
            raw = [{"text": parsed.text[start:start + size], "start": start, "end": min(start + size, len(parsed.text))} for start in range(0, len(parsed.text), size - overlap)]
        elif strategy in {"Paragraphs", "Pages"}:
            if strategy == "Pages":
                page_numbers = list(dict.fromkeys(block.page for block in parsed.blocks))
                groups = [[block for block in parsed.blocks if block.page == number] for number in page_numbers]
            else:
                groups = [[block] for block in parsed.blocks]
            for i in range(0, len(groups), size - overlap):
                selected = [block for group in groups[i:i + size] for block in group]
                start, end = selected[0].start, selected[-1].end
                raw.append({"text": parsed.text[start:end], "start": start, "end": end})
        else:
            matches = list(re.finditer(r"^#{1,6}\s+.+$", parsed.text, re.M))
            points = sorted({0, *(match.start() for match in matches), len(parsed.text)})
            raw = [{"text": parsed.text[a:b].strip()} for a, b in zip(points, points[1:])]
    elif tool == "Docling":
        if parsed.parser != "Docling":
            raise ValueError("Docling chunking requires the native Docling parser.")
        from docling_core.transforms.chunker import HierarchicalChunker, HybridChunker
        if strategy == "Hybrid":
            from docling_core.transforms.chunker.tokenizer.openai import OpenAITokenizer
            chunker = HybridChunker(tokenizer=OpenAITokenizer(tokenizer=encoding(), max_tokens=size), merge_peers=True)
        else:
            chunker = HierarchicalChunker()
        for item in chunker.chunk(dl_doc=parsed.native):
            pages = sorted({prov.page_no for doc_item in item.meta.doc_items for prov in doc_item.prov}) if Path(parsed.name).suffix.lower() == ".pdf" else []
            raw.append({"text": item.text, "pages": pages, "metadata": {"headings": item.meta.headings, "captions": item.meta.captions, "contextualized_text": chunker.contextualize(item), "elements": [str(doc_item.label) for doc_item in item.meta.doc_items]}})
    elif tool == "Unstructured":
        if parsed.parser != "Unstructured":
            raise ValueError("Unstructured chunking requires native partitioning.")
        from unstructured.chunking.basic import chunk_elements
        from unstructured.chunking.title import chunk_by_title
        kwargs = dict(max_characters=size, new_after_n_chars=size, overlap=overlap, include_orig_elements=True)
        if strategy == "Basic elements":
            items = chunk_elements(parsed.native, **kwargs)
        else:
            items = chunk_by_title(parsed.native, multipage_sections=strategy != "By title + page boundaries", combine_text_under_n_chars=0, **kwargs)
        for item in items:
            originals = getattr(item.metadata, "orig_elements", []) or []
            pages = sorted({element.metadata.page_number for element in originals if element.metadata.page_number is not None})
            raw.append({"text": str(item), "pages": pages, "metadata": {"category": item.category, "original_element_types": [element.category for element in originals]}})
    else:
        if keep_pages:
            for page in dict.fromkeys(block.page for block in parsed.blocks):
                selected = [block for block in parsed.blocks if block.page == page]
                values = text_chunks(tool, strategy, parsed.text[selected[0].start:selected[-1].end], size, overlap)
                # Convert native segment-local offsets into full-source offsets.
                for value in values:
                    if value.get("start") is not None:
                        value["start"] += selected[0].start
                        value["end"] += selected[0].start
                    value["pages"] = [page] if page is not None else []
                raw.extend(values)
        else:
            raw = text_chunks(tool, strategy, parsed.text, size, overlap)
    if tool == "Baseline" and keep_pages and strategy not in {"Pages", "Paragraphs"}:
        notes.append("Baseline character/heading strategies run on the whole text; choose Pages to enforce page boundaries.")
    if tool == "LlamaIndex" and strategy == "Hierarchical":
        notes.append("Parent and child chunks both appear. Their source text overlaps by design, so the chunk count is not directly comparable to flat splitters.")
    chunks = attach_provenance(parsed, raw)
    if units == "tokens" and not (tool == "LlamaIndex" and strategy == "Hierarchical"):
        oversized = sum(chunk.tokens > size for chunk in chunks)
        if oversized:
            notes.append(f"{oversized} native chunks exceed the requested size when recounted with cl100k_base. Native results are preserved; the library may use a different tokenization or boundary rule.")
    return {"tool": tool, "strategy": strategy, "size": size if units != "none" else None, "size_unit": units, "overlap": overlap, "keep_pages": keep_pages, "parser": parsed.parser, "page_basis": parsed.page_basis, "parse_seconds": parsed.elapsed, "chunk_seconds": round(time.perf_counter() - started, 3), "notes": notes, "chunks": [asdict(chunk) for chunk in chunks]}
